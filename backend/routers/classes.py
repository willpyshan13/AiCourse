from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from sqlalchemy.orm import Session
from typing import List
import openpyxl
import io
import datetime
import zipfile
from sqlalchemy.orm import selectinload
from config import MAX_EXCEL_BYTES
from database import get_db
from models import Class_, Team, Student
from schemas import ClassCreate, ClassInfo, TeamCreate, TeamInfo, StudentCreate, StudentInfo, ExcelImportResult
from auth import get_current_teacher
from ownership import owned_class

router = APIRouter(prefix="/api/classes", tags=["班级管理"])


def _get_teacher_id(payload: dict) -> int:
    return int(payload["sub"])


@router.get("", response_model=List[ClassInfo])
def list_classes(
    payload: dict = Depends(get_current_teacher),
    db: Session = Depends(get_db),
):
    teacher_id = _get_teacher_id(payload)
    classes = db.query(Class_).options(selectinload(Class_.teams), selectinload(Class_.students)).filter(Class_.teacher_id == teacher_id).all()
    result = []
    for c in classes:
        result.append(ClassInfo(
            id=c.id,
            teacher_id=c.teacher_id,
            name=c.name,
            grade=c.grade,
            description=c.description,
            team_count=len(c.teams),
            student_count=len(c.students),
            created_at=c.created_at,
        ))
    return result


@router.post("", response_model=ClassInfo)
def create_class(
    req: ClassCreate,
    payload: dict = Depends(get_current_teacher),
    db: Session = Depends(get_db),
):
    teacher_id = _get_teacher_id(payload)
    class_ = Class_(teacher_id=teacher_id, name=req.name, grade=req.grade, description=req.description)
    db.add(class_)
    db.commit()
    db.refresh(class_)
    return ClassInfo(
        id=class_.id,
        teacher_id=class_.teacher_id,
        name=class_.name,
        grade=class_.grade,
        description=class_.description,
        team_count=0,
        student_count=0,
        created_at=class_.created_at,
    )


@router.get("/{class_id}", response_model=ClassInfo)
def get_class(
    class_id: int,
    payload: dict = Depends(get_current_teacher),
    db: Session = Depends(get_db),
):
    class_ = owned_class(db, class_id, _get_teacher_id(payload))
    if not class_:
        raise HTTPException(status_code=404, detail="班级不存在")
    return ClassInfo(
        id=class_.id,
        teacher_id=class_.teacher_id,
        name=class_.name,
        grade=class_.grade,
        description=class_.description,
        team_count=len(class_.teams),
        student_count=len(class_.students),
        created_at=class_.created_at,
    )


@router.delete("/{class_id}")
def delete_class(
    class_id: int,
    payload: dict = Depends(get_current_teacher),
    db: Session = Depends(get_db),
):
    class_ = owned_class(db, class_id, _get_teacher_id(payload))
    if not class_:
        raise HTTPException(status_code=404, detail="班级不存在")
    db.delete(class_)
    db.commit()
    return {"ok": True}


@router.get("/{class_id}/teams", response_model=List[TeamInfo])
def list_teams(
    class_id: int,
    payload: dict = Depends(get_current_teacher),
    db: Session = Depends(get_db),
):
    owned_class(db, class_id, _get_teacher_id(payload))
    teams = db.query(Team).options(selectinload(Team.members), selectinload(Team.practices)).filter(Team.class_id == class_id).all()
    result = []
    for t in teams:
        result.append(TeamInfo(
            id=t.id,
            class_id=t.class_id,
            name=t.name,
            qr_code=t.qr_code or "",
            member_count=len(t.members),
            practice_count=len(t.practices),
            created_at=t.created_at,
        ))
    return result


@router.post("/{class_id}/teams", response_model=TeamInfo)
def create_team(
    class_id: int,
    req: TeamCreate,
    payload: dict = Depends(get_current_teacher),
    db: Session = Depends(get_db),
):
    class_ = owned_class(db, class_id, _get_teacher_id(payload))
    if not class_:
        raise HTTPException(status_code=404, detail="班级不存在")
    team = Team(class_id=class_id, name=req.name)
    db.add(team)
    db.commit()
    db.refresh(team)
    return TeamInfo(
        id=team.id,
        class_id=team.class_id,
        name=team.name,
        qr_code=team.qr_code or "",
        member_count=0,
        practice_count=0,
        created_at=team.created_at,
    )


@router.get("/{class_id}/students", response_model=List[StudentInfo])
def list_students(
    class_id: int,
    payload: dict = Depends(get_current_teacher),
    db: Session = Depends(get_db),
):
    owned_class(db, class_id, _get_teacher_id(payload))
    students = db.query(Student).filter(Student.class_id == class_id).all()
    return students


@router.post("/{class_id}/students", response_model=StudentInfo)
def create_student(
    class_id: int,
    req: StudentCreate,
    payload: dict = Depends(get_current_teacher),
    db: Session = Depends(get_db),
):
    owned_class(db, class_id, _get_teacher_id(payload))
    student = Student(
        class_id=class_id,
        name=req.name,
        student_no=req.student_no.strip(),
        phone=req.phone.strip(),
        grade=req.grade,
        birth_date=req.birth_date,
        role=req.role,
    )
    db.add(student)
    db.commit()
    db.refresh(student)
    return student


@router.post("/{class_id}/import-excel", response_model=ExcelImportResult)
def import_excel(
    class_id: int,
    file: UploadFile = File(...),
    payload: dict = Depends(get_current_teacher),
    db: Session = Depends(get_db),
):
    """
    从 Excel 导入学生数据。
    期望表头: 姓名, 学号, 年级, 出生年月, 小组, 职责, 手机号
    """
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="仅支持 .xlsx 文件")
    contents = file.file.read(MAX_EXCEL_BYTES + 1)
    if len(contents) > MAX_EXCEL_BYTES:
        raise HTTPException(status_code=413, detail="Excel 文件不能超过5MB")
    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents), read_only=True, data_only=True)
    except (OSError, ValueError, zipfile.BadZipFile, openpyxl.utils.exceptions.InvalidFileException):
        raise HTTPException(status_code=400, detail="Excel 文件无效") from None
    ws = wb.active
    if ws.max_row and ws.max_row > 5001:
        wb.close()
        raise HTTPException(status_code=400, detail="最多导入5000名学生")
    owned_class(db, class_id, _get_teacher_id(payload))

    students_imported = 0
    teams_created = 0
    errors = []

    team_cache = {team.name: team.id for team in db.query(Team).filter(Team.class_id == class_id).all()}

    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if i > 5001:
            db.rollback()
            wb.close()
            raise HTTPException(status_code=400, detail="最多导入5000名学生")
        if i == 1:
            continue  # skip header
        if not row[0]:
            continue

        try:
            name = str(row[0]).strip() if row[0] else ""
            student_no = str(row[1]).strip() if len(row) > 1 and row[1] else ""
            grade = str(row[2]).strip() if len(row) > 2 and row[2] else ""
            birth = row[3] if len(row) > 3 else None
            team_name = str(row[4]).strip() if len(row) > 4 and row[4] else ""
            role = str(row[5]).strip() if len(row) > 5 and row[5] else ""
            phone = str(row[6]).strip() if len(row) > 6 and row[6] else ""

            if not name:
                continue

            birth_date = None
            if isinstance(birth, datetime.datetime):
                birth_date = birth.date()
            elif isinstance(birth, datetime.date):
                birth_date = birth
            elif isinstance(birth, str) and birth.strip():
                try:
                    birth_date = datetime.datetime.strptime(birth.strip(), "%Y-%m-%d").date()
                except ValueError:
                    pass

            team_id = None
            if team_name:
                if team_name in team_cache:
                    team_id = team_cache[team_name]
                else:
                    team = Team(class_id=class_id, name=team_name)
                    db.add(team)
                    db.flush()
                    team_id = team.id
                    team_cache[team_name] = team_id
                    teams_created += 1

            student = Student(
                class_id=class_id,
                team_id=team_id,
                name=name,
                student_no=student_no,
                phone=phone,
                grade=grade,
                birth_date=birth_date,
                role=role,
            )
            db.add(student)
            students_imported += 1

        except Exception as e:
            errors.append(f"第{i}行: {str(e)}")

    db.commit()
    wb.close()
    return ExcelImportResult(students_imported=students_imported, teams_created=teams_created, errors=errors)
